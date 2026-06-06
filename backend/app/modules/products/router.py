"""HTTP router for products — public catalog and admin CRUD."""

import uuid
from typing import Annotated

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.pagination import PaginatedResponse
from app.dependencies import get_current_admin, get_db
from app.modules.products import service
from app.modules.products.models import GeneroEnum, Imagen, Perfume
from app.modules.products.schemas import (
    PerfumeCreate,
    PerfumeDetailResponse,
    PerfumeSummaryResponse,
    PerfumeUpdate,
    ProductFilters,
    PresentacionUpdate,
)

router = APIRouter(tags=["products"])

PREFIX_PUBLIC = "/products"
PREFIX_ADMIN = "/admin/products"


# ─── Public Catalog ─────────────────────────────────────────────────────────

@router.get(PREFIX_PUBLIC, response_model=PaginatedResponse[PerfumeSummaryResponse])
async def list_products(
    category: str | None = None,
    brand: str | None = None,
    gender: GeneroEnum | None = None,
    is_arab: bool | None = None,
    q: str | None = None,
    sort_by: str | None = "newest",
    page: Annotated[int, Query(ge=1)] = 1,
    size: Annotated[int, Query(ge=1, le=100)] = 20,
    db: AsyncSession = Depends(get_db),
):
    """List active products with optional filters. Public endpoint."""
    filters = ProductFilters(
        category=category,
        brand=brand,
        gender=gender,
        is_arab=is_arab,
        q=q,
        sort_by=sort_by,
        page=page,
        size=size,
    )
    items, total = await service.list_products(db, filters)
    return PaginatedResponse.create(
        items=items, total=total, page=page, size=size
    )


@router.get(f"{PREFIX_PUBLIC}/featured", response_model=list[PerfumeSummaryResponse])
async def get_featured(db: AsyncSession = Depends(get_db)):
    """Return featured products (destacado=True). Public."""
    return await service.get_featured_products(db)


@router.get(f"{PREFIX_PUBLIC}/newest", response_model=list[PerfumeSummaryResponse])
async def get_newest(db: AsyncSession = Depends(get_db)):
    """Return newest arrivals (es_nuevo=True). Public."""
    return await service.get_newest_products(db)



@router.get(f"{PREFIX_PUBLIC}/notes", response_model=list[dict])
async def list_notes(db: AsyncSession = Depends(get_db)):
    """Return all olfactive notes for controlled-vocabulary selectors. Public."""
    from app.modules.products.models import NotaOlfativa
    result = await db.execute(
        select(NotaOlfativa).order_by(NotaOlfativa.familia.asc(), NotaOlfativa.nombre.asc())
    )
    notes = result.scalars().all()
    return [
        {"id": str(n.id), "nombre": n.nombre, "familia": n.familia.value}
        for n in notes
    ]


@router.get(
    f"{PREFIX_PUBLIC}/{{slug}}", response_model=PerfumeDetailResponse
)
async def get_product_detail(slug: str, db: AsyncSession = Depends(get_db)):
    """Return full product detail by slug. Public."""
    return await service.get_product_by_slug(db, slug)


# ─── Admin CRUD ────────────────────────────────────────────────────────────

@router.get(
    PREFIX_ADMIN,
    response_model=PaginatedResponse[PerfumeSummaryResponse],
    dependencies=[Depends(get_current_admin)],
)
async def admin_list_products(
    q: str | None = None,
    activo: bool | None = None,
    page: Annotated[int, Query(ge=1)] = 1,
    size: Annotated[int, Query(ge=1, le=100)] = 20,
    db: AsyncSession = Depends(get_db),
):
    """
    List ALL products (active + inactive) for the admin panel.
    Supports search (q) and active filter. Admin only.
    """
    items, total = await service.admin_list_products(db, q=q, activo=activo, page=page, size=size)
    return PaginatedResponse.create(items=items, total=total, page=page, size=size)


@router.post(
    PREFIX_ADMIN,
    response_model=PerfumeDetailResponse,
    status_code=201,
    dependencies=[Depends(get_current_admin)],
)
async def create_product(
    data: PerfumeCreate, db: AsyncSession = Depends(get_db)
):
    """Create a new product with presentations and notes. Admin only."""
    return await service.create_product(db, data)


@router.get(
    f"{PREFIX_ADMIN}/{{product_id}}",
    response_model=PerfumeDetailResponse,
    dependencies=[Depends(get_current_admin)],
)
async def admin_get_product(
    product_id: uuid.UUID, db: AsyncSession = Depends(get_db)
):
    """Get full product detail by UUID. Admin only (includes inactive)."""
    return await service.admin_get_product(db, product_id)



@router.put(
    f"{PREFIX_ADMIN}/{{product_id}}",
    response_model=PerfumeDetailResponse,
    dependencies=[Depends(get_current_admin)],
)
async def update_product(
    product_id: uuid.UUID, data: PerfumeUpdate, db: AsyncSession = Depends(get_db)
):
    """Update a product's metadata. Admin only."""
    return await service.update_product(db, product_id, data)


@router.patch(
    f"{PREFIX_ADMIN}/{{product_id}}/presentations/{{presentation_id}}",
    response_model=None,
    dependencies=[Depends(get_current_admin)],
)
async def update_presentation(
    product_id: uuid.UUID,
    presentation_id: uuid.UUID,
    data: PresentacionUpdate,
    db: AsyncSession = Depends(get_db),
):
    """Inline update of a presentation's price, stock, or size. Admin only."""
    pres = await service.update_presentation(db, product_id, presentation_id, data)
    return {"id": str(pres.id), "tamano_ml": pres.tamano_ml, "precio": float(pres.precio), "stock": pres.stock}


@router.delete(
    f"{PREFIX_ADMIN}/{{product_id}}",
    status_code=204,
    dependencies=[Depends(get_current_admin)],
)
async def delete_product(
    product_id: uuid.UUID, db: AsyncSession = Depends(get_db)
):
    """Permanently delete a product. Admin only."""
    await service.delete_product(db, product_id)


@router.post(
    f"{PREFIX_ADMIN}/bulk-delete",
    dependencies=[Depends(get_current_admin)],
)
async def bulk_delete_products(
    product_ids: list[uuid.UUID],
    db: AsyncSession = Depends(get_db),
):
    """Permanently delete multiple products at once. Admin only."""
    return await service.bulk_delete_products(db, product_ids)


@router.post(
    f"{PREFIX_ADMIN}/import",
    dependencies=[Depends(get_current_admin)],
)
async def import_products(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    """
    Import products from an Excel (.xlsx) or CSV (.csv) file.

    Expected columns: nombre, marca, genero, tamano_ml, precio, stock, descripcion
    Only 'nombre' is required. Uses upsert logic: update if exists, create if not.
    Admin only.
    """
    import io
    import csv

    content = await file.read()
    filename = file.filename or ""
    rows: list[dict] = []

    if filename.endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
    elif filename.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        headers = [str(cell.value).strip().lower() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            rows.append(dict(zip(headers, row)))
    else:
        raise HTTPException(status_code=400, detail="Formato no soportado. Use .xlsx o .csv")

    result = await service.import_products_from_rows(db, rows)
    return result

